import csv
import hashlib
import io
import json
import re
from base64 import b64decode, b64encode
from datetime import datetime, timezone
from mimetypes import guess_type
from pathlib import Path
from typing import Annotated, Any
from urllib.parse import unquote_to_bytes, urlparse
from urllib.request import Request, urlopen
from uuid import uuid4

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Query, status
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from PIL import UnidentifiedImageError
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, require_roles
from app.core.audit import log_audit
from app.db.session import get_db
from app.models.ai_correction import AICorrection
from app.models.catalog_template import CatalogTemplate
from app.models.company import Company
from app.models.company_settings import CompanySettings
from app.models.image_label import ImageLabel
from app.models.marketplace_document_template import MarketplaceDocumentTemplate
from app.models.marketplace_export import MarketplaceExport
from app.models.processing_job import ProcessingJob
from app.models.product import Product
from app.models.product_image import ProductImage
from app.models.product_measurement import ProductMeasurement
from app.models.user import User
from app.schemas.catalog import (
    AICorrectionResponse,
    AnalyzeImageRequest,
    AnalyzeImageResponse,
    CatalogTemplateCreateRequest,
    CatalogTemplateListResponse,
    CatalogTemplateResponse,
    CatalogTemplateUpdateRequest,
    GenerateStyleCodeRequest,
    GenerateStyleCodeResponse,
    ImageLabelCreateRequest,
    ImageLabelListResponse,
    ImageLabelResponse,
    ImageLabelUpdateRequest,
    JobStatus,
    JobType,
    LearningFieldAccuracy,
    LearningStatsResponse,
    LogCorrectionRequest,
    MarketplaceExportCreateRequest,
    MarketplaceExportListResponse,
    MarketplaceExportResponse,
    ProcessingJobCreateRequest,
    ProcessingJobListResponse,
    ProcessingJobResponse,
    ProductCreateRequest,
    ProductImageCreateRequest,
    ProductImageResponse,
    ProductListResponse,
    ProductMeasurementCreateRequest,
    ProductMeasurementResponse,
    ProductResponse,
    ProductStatus,
    ProductUpdateRequest,
)
from app.services.ai import (
    process_ai_correction_retraining,
    process_image_analysis_job,
    process_techpack_ocr_job,
)
from app.services.marketplace_document_templates import (
    loads_json as loads_marketplace_json,
    marketplace_label,
    normalize_marketplace_key,
    normalize_template_columns,
)

router = APIRouter(prefix='/catalog', tags=['catalog'])

DbSession = Annotated[Session, Depends(get_db)]
ReadUser = Annotated[User, Depends(get_current_user)]
WriteUser = Annotated[User, Depends(require_roles('admin', 'manager', 'operator'))]

ALLOWED_EXPORT_STATUSES = {'queued', 'processing', 'completed', 'failed'}
ALLOWED_PRODUCT_STATUSES = {'draft', 'processing', 'needs_review', 'ready', 'archived'}
EXPORT_DIR = Path('static/exports')
EXPORT_DIR.mkdir(parents=True, exist_ok=True)
MAX_EXPORT_VALIDATION_ERRORS = 20
AI_ANALYSIS_FIELDS = ('category', 'style_name', 'color', 'fabric', 'composition', 'woven_knits')
ANALYZE_ALLOWED_FIELD_MAP = {
    'category': 'allowed_categories',
    'style_name': 'allowed_style_names',
    'color': 'allowed_colors',
    'fabric': 'allowed_fabrics',
    'composition': 'allowed_compositions',
    'woven_knits': 'allowed_woven_knits',
}
MAX_AI_ALLOWED_OPTIONS_PER_FIELD = 40
MAX_AI_CORRECTION_HINTS = 36
API_ROOT_DIR = Path(__file__).resolve().parents[4]
API_STATIC_DIR = API_ROOT_DIR / 'static'

MARKETPLACE_ALIASES = {
    'myntra': 'myntra',
    'ajio': 'ajio',
    'amazon': 'amazon_in',
    'amazon in': 'amazon_in',
    'amazon india': 'amazon_in',
    'amazon_in': 'amazon_in',
    'flipkart': 'flipkart',
    'nykaa': 'nykaa',
    'generic': 'generic',
}

MARKETPLACE_EXPORT_TEMPLATES: dict[str, dict[str, Any]] = {
    'generic': {
        'label': 'Generic',
        'required_fields': ('sku', 'title'),
        'columns': (
            ('S. No', 'serial_no'),
            ('Style-No', 'sku'),
            ('Image Preview', 'image_preview'),
            ('Name', 'title'),
            ('Category', 'category'),
            ('Color', 'color'),
            ('Fabric', 'fabric'),
            ('Composition', 'composition'),
            ('Woven/Knits', 'woven_knits'),
            ('Units', 'units'),
            ('PO Price', 'po_price'),
            ('OSP', 'osp'),
            ('Status', 'status'),
        ),
    },
    'myntra': {
        'label': 'Myntra',
        'required_fields': ('sku', 'title', 'brand', 'category', 'color', 'size', 'mrp'),
        'columns': (
            ('S. No', 'serial_no'),
            ('Style ID', 'sku'),
            ('Image Preview', 'image_preview'),
            ('Product Name', 'title'),
            ('Brand', 'brand'),
            ('Category', 'category'),
            ('Color', 'color'),
            ('Size', 'size'),
            ('MRP', 'mrp'),
            ('Description', 'description'),
        ),
    },
    'ajio': {
        'label': 'Ajio',
        'required_fields': ('sku', 'title', 'category', 'color', 'size', 'mrp'),
        'columns': (
            ('S. No', 'serial_no'),
            ('Seller SKU', 'sku'),
            ('Image Preview', 'image_preview'),
            ('Product Name', 'title'),
            ('Department', 'category'),
            ('Color', 'color'),
            ('Size', 'size'),
            ('MRP', 'mrp'),
            ('Description', 'description'),
        ),
    },
    'amazon_in': {
        'label': 'Amazon IN',
        'required_fields': ('sku', 'title', 'brand', 'category', 'mrp'),
        'columns': (
            ('S. No', 'serial_no'),
            ('seller-sku', 'sku'),
            ('image-preview', 'image_preview'),
            ('item-name', 'title'),
            ('brand-name', 'brand'),
            ('item-type', 'category'),
            ('color-name', 'color'),
            ('size-name', 'size'),
            ('standard-price', 'mrp'),
            ('product-description', 'description'),
        ),
    },
    'flipkart': {
        'label': 'Flipkart',
        'required_fields': ('sku', 'title', 'brand', 'category', 'color', 'size', 'mrp'),
        'columns': (
            ('S. No', 'serial_no'),
            ('Seller SKU', 'sku'),
            ('Image Preview', 'image_preview'),
            ('Product Title', 'title'),
            ('Brand', 'brand'),
            ('Category', 'category'),
            ('Color', 'color'),
            ('Size', 'size'),
            ('Selling Price', 'mrp'),
            ('Description', 'description'),
        ),
    },
    'nykaa': {
        'label': 'Nykaa',
        'required_fields': ('sku', 'title', 'brand', 'category', 'color', 'mrp'),
        'columns': (
            ('S. No', 'serial_no'),
            ('SKU', 'sku'),
            ('Image Preview', 'image_preview'),
            ('Name', 'title'),
            ('Brand', 'brand'),
            ('Category', 'category'),
            ('Shade', 'color'),
            ('Size', 'size'),
            ('MRP', 'mrp'),
            ('Description', 'description'),
        ),
    },
}


def _json_loads(raw: str | None) -> dict[str, Any]:
    if not raw:
        return {}
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _json_dumps(payload: dict[str, Any] | None) -> str:
    return json.dumps(payload or {}, separators=(',', ':'))


def _json_list_loads(raw: str | None) -> list[str]:
    if not raw:
        return []
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return []
    if not isinstance(parsed, list):
        return []
    cleaned: list[str] = []
    for item in parsed:
        if not isinstance(item, str):
            continue
        token = item.strip()
        if token:
            cleaned.append(token)
    return cleaned


def _normalize_template_tokens(values: list[str] | None) -> list[str]:
    if not values:
        return []
    cleaned = [value.strip() for value in values if isinstance(value, str) and value.strip()]
    return list(dict.fromkeys(cleaned))


def _template_defaults_without_meta(defaults: dict[str, Any]) -> dict[str, Any]:
    cleaned = dict(defaults)
    cleaned.pop('_allowed', None)
    return cleaned


def _template_allowed_from_defaults(defaults: dict[str, Any]) -> dict[str, list[str]]:
    raw_allowed = defaults.get('_allowed')
    if not isinstance(raw_allowed, dict):
        return {
            'categories': [],
            'style_names': [],
            'compositions': [],
            'woven_knits': [],
        }
    return {
        'categories': _normalize_template_tokens(raw_allowed.get('categories') if isinstance(raw_allowed.get('categories'), list) else []),
        'style_names': _normalize_template_tokens(raw_allowed.get('style_names') if isinstance(raw_allowed.get('style_names'), list) else []),
        'compositions': _normalize_template_tokens(raw_allowed.get('compositions') if isinstance(raw_allowed.get('compositions'), list) else []),
        'woven_knits': _normalize_template_tokens(raw_allowed.get('woven_knits') if isinstance(raw_allowed.get('woven_knits'), list) else []),
    }


def _merge_template_defaults_with_allowed(
    defaults: dict[str, Any] | None,
    *,
    allowed_categories: list[str],
    allowed_style_names: list[str],
    allowed_compositions: list[str],
    allowed_woven_knits: list[str],
) -> dict[str, Any]:
    payload = _template_defaults_without_meta(defaults or {})
    payload['_allowed'] = {
        'categories': _normalize_template_tokens(allowed_categories),
        'style_names': _normalize_template_tokens(allowed_style_names),
        'compositions': _normalize_template_tokens(allowed_compositions),
        'woven_knits': _normalize_template_tokens(allowed_woven_knits),
    }
    return payload


def _normalize_marketplace(raw_marketplace: str) -> str:
    try:
        return normalize_marketplace_key(raw_marketplace)
    except ValueError as err:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(err)) from err


def _normalize_export_filters(raw_filters: dict[str, Any] | None) -> dict[str, Any]:
    filters: dict[str, Any] = {}
    if not raw_filters:
        return filters

    if 'status' in raw_filters:
        raw_status = raw_filters['status']
        if isinstance(raw_status, str):
            cleaned = raw_status.strip()
            if cleaned not in ALLOWED_PRODUCT_STATUSES:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f'Unsupported product status filter "{cleaned}".',
                )
            filters['status'] = cleaned
        elif isinstance(raw_status, list):
            cleaned_statuses = [item.strip() for item in raw_status if isinstance(item, str) and item.strip()]
            invalid = [item for item in cleaned_statuses if item not in ALLOWED_PRODUCT_STATUSES]
            if invalid:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f'Unsupported product status filters: {", ".join(invalid)}.',
                )
            if cleaned_statuses:
                filters['status'] = cleaned_statuses
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail='The "status" filter must be a string or array of strings.',
            )
    if 'search' in raw_filters and isinstance(raw_filters['search'], str):
        search = raw_filters['search'].strip()
        if search:
            filters['search'] = search[:120]

    if 'product_ids' in raw_filters:
        raw_product_ids = raw_filters['product_ids']
        if not isinstance(raw_product_ids, list):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail='The "product_ids" filter must be an array of product ids.',
            )
        product_ids = [str(value).strip() for value in raw_product_ids if str(value).strip()]
        if product_ids:
            filters['product_ids'] = list(dict.fromkeys(product_ids))

    return filters


def _render_pattern_sku(
    db: Session,
    company_id: str,
    pattern: str,
    payload: GenerateStyleCodeRequest,
) -> str:
    token_map = _build_sku_token_map(
        brand=payload.brand,
        category=payload.category,
        color=None,
        size=None,
    )
    rendered = pattern.upper()
    for key, value in token_map.items():
        rendered = rendered.replace(key, value)
    rendered = re.sub(r'[^A-Z0-9-]+', '-', rendered)
    rendered = re.sub(r'-{2,}', '-', rendered).strip('-')
    if not rendered:
        rendered = 'SKU'
    return _find_unique_sku(db, company_id, rendered)


def _build_sku_token_map(
    *,
    brand: str | None,
    category: str | None,
    color: str | None,
    size: str | None,
) -> dict[str, str]:
    current_year = str(datetime.now(timezone.utc).year)
    return {
        '{BRAND}': _sku_token(brand, 'GEN'),
        '{CATEGORY}': _sku_token(category, 'CAT'),
        '{COLOR}': _sku_token(color, 'NA'),
        '{SIZE}': _sku_token(size, 'OS'),
        '{YEAR}': current_year,
        '{YY}': current_year[-2:],
    }


def _render_sku_from_format(
    *,
    sku_format: str,
    brand: str | None,
    category: str | None,
    color: str | None,
    size: str | None,
) -> str:
    rendered = sku_format.upper()
    for key, value in _build_sku_token_map(
        brand=brand,
        category=category,
        color=color,
        size=size,
    ).items():
        rendered = rendered.replace(key, value)

    rendered = re.sub(r'[^A-Z0-9-]+', '-', rendered)
    rendered = re.sub(r'-{2,}', '-', rendered).strip('-')
    return rendered or 'SKU'


def _compute_image_hash(image_url: str) -> str:
    normalized = image_url.strip()
    try:
        if normalized.startswith('data:'):
            header, _, payload = normalized.partition(',')
            if ';base64' in header:
                image_bytes = b64decode(payload.encode('utf-8'), validate=False)
            else:
                image_bytes = unquote_to_bytes(payload)
        else:
            image_bytes = normalized.encode('utf-8')
    except Exception:
        image_bytes = normalized.encode('utf-8')
    return hashlib.sha256(image_bytes).hexdigest()


def _coerce_local_image_url_to_data_url(image_url: str) -> str:
    normalized = image_url.strip()
    if not normalized or normalized.startswith('data:'):
        return normalized

    parsed = urlparse(normalized)
    local_static_paths: list[Path] = []

    # Handle both absolute URLs (/static/...) and full URLs (.../static/...).
    static_relative = ''
    if normalized.startswith('/static/'):
        static_relative = normalized.removeprefix('/static/')
    elif parsed.path.startswith('/static/'):
        static_relative = parsed.path.removeprefix('/static/')

    if static_relative:
        local_static_paths.append(API_STATIC_DIR / static_relative)
        local_static_paths.append(Path('static') / static_relative)

    for local_path in local_static_paths:
        if local_path.exists():
            image_bytes = local_path.read_bytes()
            mime_type = guess_type(local_path.name)[0] or 'image/jpeg'
            encoded = b64encode(image_bytes).decode('ascii')
            return f'data:{mime_type};base64,{encoded}'

    # If image is referenced via localhost, fetch it server-side and convert.
    if parsed.scheme in {'http', 'https'} and parsed.hostname in {'127.0.0.1', 'localhost'}:
        try:
            request = Request(normalized, headers={'User-Agent': 'kira-ai-analyzer/1.0'})
            with urlopen(request, timeout=6) as response:
                image_bytes = response.read()
                mime_type = response.headers.get_content_type() or 'image/jpeg'
            if image_bytes:
                encoded = b64encode(image_bytes).decode('ascii')
                return f'data:{mime_type};base64,{encoded}'
        except Exception:
            return normalized

    return normalized


def _normalize_analysis_field(raw_field: Any) -> dict[str, Any]:
    if isinstance(raw_field, dict):
        raw_value = raw_field.get('value')
        raw_confidence = raw_field.get('confidence')
        source = raw_field.get('source')
        based_on = raw_field.get('based_on') or raw_field.get('basedOn')
        learned_from = raw_field.get('learned_from') or raw_field.get('learnedFrom')
    else:
        raw_value = raw_field
        raw_confidence = None
        source = None
        based_on = None
        learned_from = None

    value = str(raw_value).strip() if isinstance(raw_value, str) and raw_value.strip() else None
    confidence: float | None = None
    if isinstance(raw_confidence, (int, float)):
        confidence = max(0.0, min(100.0, float(raw_confidence)))

    normalized_source = source.strip() if isinstance(source, str) and source.strip() else 'vision_model'
    normalized_based_on = (
        based_on.strip()
        if isinstance(based_on, str) and based_on.strip()
        else 'Image texture, silhouette, and color cues'
    )
    normalized_learned_from = (
        learned_from.strip()
        if isinstance(learned_from, str) and learned_from.strip()
        else 'Catalog priors and historical apparel patterns'
    )

    return {
        'value': value,
        'confidence': confidence,
        'source': normalized_source,
        'based_on': normalized_based_on,
        'learned_from': normalized_learned_from,
    }


def _normalize_analysis_result(raw_result: Any, image_hash: str) -> AnalyzeImageResponse:
    payload: dict[str, Any]
    if isinstance(raw_result, dict):
        payload = raw_result
    else:
        payload = {}

    normalized: dict[str, Any] = {field: _normalize_analysis_field(payload.get(field)) for field in AI_ANALYSIS_FIELDS}
    normalized['image_hash'] = image_hash
    return AnalyzeImageResponse(**normalized)


def _normalize_analyze_allowed_options(raw_allowed: dict[str, Any] | None) -> dict[str, list[str]]:
    if not isinstance(raw_allowed, dict):
        return {}

    normalized: dict[str, list[str]] = {}
    for field_name, payload_key in ANALYZE_ALLOWED_FIELD_MAP.items():
        raw_values = raw_allowed.get(payload_key)
        if not isinstance(raw_values, list):
            continue
        string_values = [value for value in raw_values if isinstance(value, str)]
        cleaned = _normalize_template_tokens(string_values)[:MAX_AI_ALLOWED_OPTIONS_PER_FIELD]
        if cleaned:
            normalized[field_name] = cleaned
    return normalized


def _build_recent_correction_hints(db: Session, company_id: str) -> list[dict[str, str]]:
    rows = (
        db.query(AICorrection)
        .filter(
            AICorrection.company_id == company_id,
            AICorrection.feedback_type == 'reject',
            AICorrection.corrected_value.isnot(None),
            AICorrection.field_name.in_(AI_ANALYSIS_FIELDS),
        )
        .order_by(AICorrection.created_at.desc())
        .limit(MAX_AI_CORRECTION_HINTS * 3)
        .all()
    )

    hints: list[dict[str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    for row in rows:
        field_name = (row.field_name or '').strip().lower()
        if field_name not in AI_ANALYSIS_FIELDS:
            continue
        corrected = (row.corrected_value or '').strip()
        if not corrected:
            continue
        suggested = (row.suggested_value or '').strip()
        dedupe_key = (field_name, suggested.lower(), corrected.lower())
        if dedupe_key in seen:
            continue
        seen.add(dedupe_key)
        hint = {'field': field_name, 'corrected_value': corrected}
        if suggested:
            hint['suggested_value'] = suggested
        hints.append(hint)
        if len(hints) >= MAX_AI_CORRECTION_HINTS:
            break
    return hints


def _pick_mrp(ai_attributes: dict[str, Any]) -> str:
    for key in ('mrp', 'price', 'standard_price', 'selling_price'):
        value = ai_attributes.get(key)
        if value in (None, ''):
            continue
        if isinstance(value, (int, float)):
            return f'{value:.2f}'
        return str(value).strip()
    return ''


def _pick_text(ai_attributes: dict[str, Any], keys: tuple[str, ...], fallback: str = '') -> str:
    for key in keys:
        value = ai_attributes.get(key)
        if value in (None, ''):
            continue
        return str(value).strip()
    return fallback


def _pick_osp(ai_attributes: dict[str, Any]) -> str:
    for key in ('osp', 'osp_sar', 'osp_price', 'price'):
        value = ai_attributes.get(key)
        if value in (None, ''):
            continue
        if isinstance(value, (int, float)):
            return f'SAR {value}'
        text_value = str(value).strip()
        if text_value.upper().startswith('SAR'):
            return text_value
        return f'SAR {text_value}'
    return 'SAR 95'


def _normalize_export_image_url(value: str | None) -> str:
    if value is None:
        return ''
    cleaned = value.strip()
    if not cleaned or cleaned.startswith('uploaded://'):
        return ''
    return cleaned


def _extract_export_fields(product: Product, image_url: str = '') -> dict[str, str]:
    ai_attributes = _json_loads(product.ai_attributes_json)
    default_brand = 'Generic'
    default_size = 'Free Size'
    default_mrp = '0.00'
    default_category = 'DRESSES'
    default_color = 'Blue'
    return {
        'sku': (product.sku or '').strip(),
        'title': (product.title or '').strip(),
        'brand': (product.brand or '').strip() or default_brand,
        'category': (product.category or '').strip() or default_category,
        'color': (product.color or '').strip() or default_color,
        'size': (product.size or '').strip() or default_size,
        'description': (product.description or '').strip(),
        'mrp': _pick_mrp(ai_attributes) or default_mrp,
        'fabric': _pick_text(ai_attributes, ('fabric',), fallback='Poly Georgette'),
        'composition': _pick_text(ai_attributes, ('composition',), fallback='100% Polyester'),
        'woven_knits': _pick_text(ai_attributes, ('woven_knits', 'wovenKnits'), fallback='Woven'),
        'units': _pick_text(ai_attributes, ('units', 'total_units', 'totalUnits'), fallback='24'),
        'po_price': _pick_text(ai_attributes, ('po_price', 'poPrice', 'po_price_sar'), fallback='600'),
        'osp': _pick_osp(ai_attributes),
        'status': (product.status or '').strip(),
        'image_url': _normalize_export_image_url(image_url),
        'image_preview': _normalize_export_image_url(image_url),
    }


def _catalog_builtin_template_config(marketplace_key: str) -> dict[str, Any]:
    template = MARKETPLACE_EXPORT_TEMPLATES[marketplace_key]
    required_fields = set(template['required_fields'])
    columns = [
        {'header': header, 'source_field': source_field, 'required': source_field in required_fields}
        for header, source_field in template['columns']
    ]
    return {
        'marketplace_key': marketplace_key,
        'label': template['label'],
        'columns': columns,
        'required_fields': required_fields,
    }


def _catalog_template_config_from_record(record: MarketplaceDocumentTemplate) -> dict[str, Any]:
    schema_payload = loads_marketplace_json(record.schema_json)
    columns = normalize_template_columns(schema_payload.get('columns') if isinstance(schema_payload, dict) else [])
    required_fields = {
        str(column.get('source_field'))
        for column in columns
        if column.get('source_field') and bool(column.get('required'))
    }
    if not columns:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f'Marketplace template "{record.name}" has no mapped columns.',
        )
    return {
        'marketplace_key': record.marketplace_key,
        'label': marketplace_label(record.marketplace_key),
        'columns': columns,
        'required_fields': required_fields,
    }


def _collect_export_validation_errors(
    products: list[Product], template_config: dict[str, Any], primary_image_map: dict[str, str]
) -> list[str]:
    required_fields = set(template_config['required_fields'])
    errors: list[str] = []

    for product in products:
        fields = _extract_export_fields(product, primary_image_map.get(product.id, ''))
        missing_fields = [field for field in required_fields if fields.get(field, '') == '']
        if missing_fields:
            missing_label = ', '.join(missing_fields)
            errors.append(f'{product.sku}: missing {missing_label}')
        if len(errors) >= MAX_EXPORT_VALIDATION_ERRORS:
            break
    return errors


def _query_export_products(db: Session, company_id: str, filters: dict[str, Any]) -> list[Product]:
    query = db.query(Product).filter(Product.company_id == company_id)

    status_filter = filters.get('status')
    if isinstance(status_filter, list) and status_filter:
        query = query.filter(Product.status.in_(status_filter))
    elif isinstance(status_filter, str) and status_filter:
        query = query.filter(Product.status == status_filter)

    search = filters.get('search')
    if isinstance(search, str) and search:
        term = f'%{search}%'
        query = query.filter(or_(Product.sku.ilike(term), Product.title.ilike(term), Product.category.ilike(term)))

    product_ids = filters.get('product_ids')
    if isinstance(product_ids, list) and product_ids:
        query = query.filter(Product.id.in_(product_ids))

    return query.order_by(Product.updated_at.desc()).all()


def _build_primary_image_map(db: Session, company_id: str, product_ids: list[str]) -> dict[str, str]:
    if not product_ids:
        return {}
    rows = (
        db.query(ProductImage)
        .filter(ProductImage.company_id == company_id, ProductImage.product_id.in_(product_ids))
        .order_by(ProductImage.created_at.asc())
        .all()
    )
    image_map: dict[str, str] = {}
    for image in rows:
        if image.product_id in image_map:
            continue
        image_map[image.product_id] = image.file_url
    return image_map


def _generate_csv_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue().encode('utf-8')


def _resolve_export_image_path(image_url: str) -> Path | None:
    normalized = image_url.strip()
    if not normalized:
        return None

    for prefix in ('/api/v1/static/', 'api/v1/static/'):
        if normalized.startswith(prefix):
            suffix = normalized.split('/static/', 1)[-1]
            candidate = API_STATIC_DIR / suffix
            return candidate if candidate.exists() else None

    for prefix in ('/static/', 'static/'):
        if normalized.startswith(prefix):
            suffix = normalized.split('static/', 1)[-1]
            candidate = API_STATIC_DIR / suffix
            return candidate if candidate.exists() else None

    return None


def _load_export_image_bytes(image_url: str) -> bytes | None:
    normalized = image_url.strip()
    if not normalized:
        return None

    if normalized.startswith('data:'):
        try:
            _, payload = normalized.split(',', 1)
        except ValueError:
            return None
        try:
            return b64decode(payload.encode('utf-8'), validate=False)
        except ValueError:
            try:
                return unquote_to_bytes(payload)
            except Exception:
                return None

    local_path = _resolve_export_image_path(normalized)
    if local_path is not None:
        try:
            return local_path.read_bytes()
        except OSError:
            return None

    parsed = urlparse(normalized)
    if parsed.scheme in {'http', 'https'}:
        try:
            request = Request(normalized, headers={'User-Agent': 'kira-export/1.0'})
            with urlopen(request, timeout=10) as response:
                return response.read()
        except Exception:
            return None

    return None


def _build_export_workbook_image(image_url: str) -> OpenpyxlImage | None:
    raw_bytes = _load_export_image_bytes(image_url)
    if not raw_bytes:
        return None
    try:
        with PILImage.open(io.BytesIO(raw_bytes)) as source:
            source.load()
            prepared = source.convert('RGBA') if source.mode not in {'RGB', 'RGBA'} else source.copy()
            prepared.thumbnail((96, 96))
            output = io.BytesIO()
            prepared.save(output, format='PNG')
    except (OSError, UnidentifiedImageError, ValueError):
        return None

    output.seek(0)
    image = OpenpyxlImage(output)
    image.width = min(image.width, 96)
    image.height = min(image.height, 96)
    image._data_stream = output  # keep stream alive until workbook save
    return image


def _generate_xlsx_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Export'
    sheet.append(headers)

    image_preview_column: int | None = None
    for column_index, header in enumerate(headers, start=1):
        normalized_header = header.strip().lower()
        if normalized_header in {'image preview', 'image-preview'}:
            image_preview_column = column_index
        width = 20 if normalized_header in {'image preview', 'image-preview'} else 28
        if normalized_header == 's. no':
            width = 10
        if normalized_header in {'primary image url', 'main-image-url'}:
            width = 42
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index, row in enumerate(rows, start=2):
        row_values = list(row)
        if image_preview_column is not None:
            row_values[image_preview_column - 1] = ''
        sheet.append(row_values)
        if image_preview_column is None:
            continue
        image_url = row[image_preview_column - 1]
        if not image_url:
            continue
        embedded_image = _build_export_workbook_image(image_url)
        if embedded_image is None:
            continue
        cell_ref = f'{get_column_letter(image_preview_column)}{row_index}'
        sheet.row_dimensions[row_index].height = 76
        sheet.add_image(embedded_image, cell_ref)

    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{sheet.max_row}'

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_export_file(
    products: list[Product], template_config: dict[str, Any], export_format: str, primary_image_map: dict[str, str]
) -> tuple[bytes, int]:
    columns = list(template_config['columns'])
    headers = [str(column.get('header') or '') for column in columns]
    rows: list[list[str]] = []
    for index, product in enumerate(products, start=1):
        fields = _extract_export_fields(product, primary_image_map.get(product.id, ''))
        fields['serial_no'] = str(index)
        data_row = [fields.get(str(column.get('source_field') or ''), '') for column in columns]
        rows.append(data_row)

    if export_format == 'csv':
        return _generate_csv_bytes(headers, rows), len(rows)
    return _generate_xlsx_bytes(headers, rows), len(rows)


def _write_export_file(content: bytes, marketplace_key: str, export_format: str) -> str:
    file_name = f'{marketplace_key}-{datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")}-{uuid4().hex[:8]}.{export_format}'
    file_path = EXPORT_DIR / file_name
    file_path.write_bytes(content)
    return f'/static/exports/{file_name}'


def _sku_token(value: str | None, fallback: str) -> str:
    cleaned = re.sub(r'[^A-Za-z0-9]+', '-', (value or '').strip()).strip('-').upper()
    return cleaned or fallback


def _normalize_explicit_sku(value: str) -> str:
    cleaned = re.sub(r'[^A-Za-z0-9-]+', '-', value.strip()).strip('-').upper()
    if not cleaned:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail='SKU cannot be empty.')
    return cleaned


def _find_unique_sku(db: Session, company_id: str, base_sku: str, current_product_id: str | None = None) -> str:
    candidate = base_sku
    suffix = 1
    while True:
        query = db.query(Product).filter(Product.company_id == company_id, Product.sku == candidate)
        if current_product_id is not None:
            query = query.filter(Product.id != current_product_id)
        if query.first() is None:
            return candidate
        suffix += 1
        candidate = f'{base_sku}-{suffix:02d}'


def _get_brand_initials(brand: str | None, fallback: str = 'GN') -> str:
    fallback_clean = re.sub(r'[^A-Za-z0-9]+', '', fallback).upper() or 'GN'
    if not brand:
        return (fallback_clean + 'XX')[:2]

    words = [w for w in re.split(r'[^A-Za-z0-9]+', brand) if w]
    code = ''
    if len(words) == 1:
        code = words[0][:2].upper()
    elif len(words) >= 2:
        code = ''.join(w[0] for w in words)[:2].upper()
    return (code + fallback_clean)[:2]

def _get_category_abbr(category: str | None, fallback: str = 'XX') -> str:
    fallback_clean = re.sub(r'[^A-Za-z0-9]+', '', fallback).upper() or 'XX'
    if not category:
        return (fallback_clean + 'XX')[:2]
    clean = re.sub(r'[^A-Za-z0-9]+', '', category).upper()
    if 'DRESS' in clean: return 'DS'
    if 'CORD' in clean or 'SET' in clean: return 'CS'
    if 'TOP' in clean: return 'TP'
    if len(clean) >= 2: return clean[:2]
    return (clean + fallback_clean)[:2]

def _find_formatted_unique_sku(db: Session, company_id: str, base_sku: str, current_product_id: str | None = None) -> str:
    suffix = 1
    while True:
        candidate = f'{base_sku}{suffix:03d}'
        query = db.query(Product).filter(Product.company_id == company_id, Product.sku == candidate)
        if current_product_id:
            query = query.filter(Product.id != current_product_id)
        if query.first() is None:
            return candidate
        suffix += 1


def _generate_default_catalog_style_no(db: Session, company_id: str, category: str | None) -> str:
    company = db.query(Company).filter(Company.id == company_id).first()
    company_name = company.name if company else None
    brand_code = _get_brand_initials(company_name, fallback='GN')
    category_code = _get_category_abbr(category, fallback='XX')
    current_year_short = datetime.now(timezone.utc).strftime('%y')
    base_sku = f'{brand_code}{category_code}{current_year_short}'
    return _find_formatted_unique_sku(db, company_id, base_sku)


def _generate_sku(db: Session, company_id: str, company_settings: CompanySettings | None, payload: ProductCreateRequest) -> str:
    default_format = '{BRAND}-{CATEGORY}-{COLOR}-{SIZE}'
    configured_format = (
        company_settings.sku_format.strip()
        if company_settings and company_settings.sku_format and company_settings.sku_format.strip()
        else default_format
    )

    # Backward compatibility: if tenant explicitly configured a non-default format,
    # keep honoring it.
    if configured_format != default_format:
        company = db.query(Company).filter(Company.id == company_id).first()
        company_name = company.name if company else 'Generic'
        brand_val = payload.brand
        if not brand_val or brand_val.strip().lower() == 'generic':
            brand_val = company_name

        base_sku = _render_sku_from_format(
            sku_format=configured_format,
            brand=brand_val,
            category=payload.category,
            color=payload.color,
            size=payload.size,
        )
        return _find_unique_sku(db, company_id, base_sku)

    return _generate_default_catalog_style_no(db, company_id, payload.category)


def _to_product_response(product: Product, *, primary_image_url: str | None = None) -> ProductResponse:
    return ProductResponse(
        id=product.id,
        company_id=product.company_id,
        sku=product.sku,
        title=product.title,
        description=product.description,
        brand=product.brand,
        category=product.category,
        color=product.color,
        size=product.size,
        status=product.status,
        confidence_score=float(product.confidence_score) if product.confidence_score is not None else None,
        ai_attributes=_json_loads(product.ai_attributes_json),
        created_by_user_id=product.created_by_user_id,
        created_at=product.created_at,
        updated_at=product.updated_at,
        primary_image_url=primary_image_url,
    )


def _to_image_response(image: ProductImage) -> ProductImageResponse:
    return ProductImageResponse(
        id=image.id,
        product_id=image.product_id,
        file_name=image.file_name,
        file_url=image.file_url,
        mime_type=image.mime_type,
        file_size_bytes=image.file_size_bytes,
        width_px=image.width_px,
        height_px=image.height_px,
        processing_status=image.processing_status,
        analysis=_json_loads(image.analysis_json),
        created_at=image.created_at,
    )


def _to_measurement_response(measurement: ProductMeasurement) -> ProductMeasurementResponse:
    return ProductMeasurementResponse(
        id=measurement.id,
        product_id=measurement.product_id,
        measurement_key=measurement.measurement_key,
        measurement_value=float(measurement.measurement_value),
        unit=measurement.unit,
        source=measurement.source,
        confidence_score=float(measurement.confidence_score) if measurement.confidence_score is not None else None,
        needs_review=measurement.needs_review,
        notes=measurement.notes,
        created_at=measurement.created_at,
    )


def _to_export_response(record: MarketplaceExport) -> MarketplaceExportResponse:
    return MarketplaceExportResponse(
        id=record.id,
        company_id=record.company_id,
        requested_by_user_id=record.requested_by_user_id,
        marketplace=record.marketplace,
        template_id=record.template_id,
        template_name=record.template_name,
        export_format=record.export_format,
        status=record.status,
        filters=_json_loads(record.filters_json),
        file_url=record.file_url,
        error_message=record.error_message,
        row_count=record.row_count,
        created_at=record.created_at,
        completed_at=record.completed_at,
    )


def _to_catalog_template_response(record: CatalogTemplate) -> CatalogTemplateResponse:
    defaults_payload = _json_loads(record.defaults_json)
    allowed_meta = _template_allowed_from_defaults(defaults_payload)
    return CatalogTemplateResponse(
        id=record.id,
        company_id=record.company_id,
        name=record.name,
        description=record.description,
        defaults=_template_defaults_without_meta(defaults_payload),
        allowed_categories=allowed_meta['categories'],
        allowed_style_names=allowed_meta['style_names'],
        allowed_colors=_json_list_loads(record.allowed_colors_json),
        allowed_fabrics=_json_list_loads(record.allowed_fabrics_json),
        allowed_compositions=allowed_meta['compositions'],
        allowed_woven_knits=allowed_meta['woven_knits'],
        style_code_pattern=record.style_code_pattern,
        is_active=record.is_active,
        created_by_user_id=record.created_by_user_id,
        created_at=record.created_at,
        updated_at=record.updated_at,
    )


def _to_job_response(record: ProcessingJob) -> ProcessingJobResponse:
    return ProcessingJobResponse(
        id=record.id,
        company_id=record.company_id,
        product_id=record.product_id,
        job_type=record.job_type,
        status=record.status,
        input_ref=record.input_ref,
        payload=_json_loads(record.payload_json),
        result=_json_loads(record.result_json),
        confidence_score=float(record.confidence_score) if record.confidence_score is not None else None,
        progress_percent=record.progress_percent,
        error_message=record.error_message,
        created_by_user_id=record.created_by_user_id,
        created_at=record.created_at,
        started_at=record.started_at,
        completed_at=record.completed_at,
    )


def _to_ai_correction_response(record: AICorrection) -> AICorrectionResponse:
    return AICorrectionResponse(
        id=record.id,
        company_id=record.company_id,
        product_id=record.product_id,
        image_hash=record.image_hash,
        field_name=record.field_name,
        feedback_type=record.feedback_type,
        suggested_value=record.suggested_value,
        corrected_value=record.corrected_value,
        reason_code=record.reason_code,
        notes=record.notes,
        source=record.source,
        based_on=record.based_on,
        learned_from=record.learned_from,
        confidence_score=float(record.confidence_score) if record.confidence_score is not None else None,
        retraining_status=record.retraining_status,
        retraining_notes=record.retraining_notes,
        created_by_user_id=record.created_by_user_id,
        created_at=record.created_at,
        processed_at=record.processed_at,
    )


def _normalize_optional_label_text(value: str | None) -> str | None:
    if value is None:
        return None
    cleaned = value.strip()
    return cleaned or None


def _label_value_changed(ai_value: str | None, human_value: str | None) -> bool:
    if not ai_value or not human_value:
        return False
    return ai_value.strip().lower() != human_value.strip().lower()


def _compute_label_corrected(
    ai_category: str | None,
    ai_style: str | None,
    human_category: str | None,
    human_style: str | None,
) -> bool:
    return _label_value_changed(ai_category, human_category) or _label_value_changed(ai_style, human_style)


def _to_image_label_response(record: ImageLabel) -> ImageLabelResponse:
    return ImageLabelResponse(
        id=record.id,
        company_id=record.company_id,
        image_url=record.image_url,
        ai_category=record.ai_category,
        ai_style=record.ai_style,
        human_category=record.human_category,
        human_style=record.human_style,
        corrected=record.corrected,
        created_at=record.created_at,
    )


def _build_learning_insights(items_processed: int, field_accuracy: list[LearningFieldAccuracy]) -> list[str]:
    insights: list[str] = []
    if items_processed > 0:
        insights.append(f'AI analyzed {items_processed} catalog image(s) so far.')
    if field_accuracy:
        top = max(field_accuracy, key=lambda item: item.accuracy_percent)
        weakest = min(field_accuracy, key=lambda item: item.accuracy_percent)
        insights.append(
            f'Highest confidence alignment is on {top.field_name.replace("_", " ").title()} at {top.accuracy_percent:.1f}%.'
        )
        if weakest.field_name != top.field_name:
            insights.append(
                f'Most corrections are needed for {weakest.field_name.replace("_", " ").title()} ({weakest.accuracy_percent:.1f}% accepted).'
            )
    if not insights:
        insights.append('No AI learning data yet. Start by running image analysis and giving field feedback.')
    return insights


def _get_company_product_or_404(db: Session, company_id: str, product_id: str) -> Product:
    product = db.query(Product).filter(Product.id == product_id, Product.company_id == company_id).first()
    if product is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Product not found.')
    return product


def _get_company_template_or_404(db: Session, company_id: str, template_id: str) -> CatalogTemplate:
    template = (
        db.query(CatalogTemplate)
        .filter(CatalogTemplate.id == template_id, CatalogTemplate.company_id == company_id)
        .first()
    )
    if template is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Template not found.')
    return template


def _get_marketplace_document_template_or_404(
    db: Session, company_id: str, template_id: str
) -> MarketplaceDocumentTemplate:
    template = (
        db.query(MarketplaceDocumentTemplate)
        .filter(
            MarketplaceDocumentTemplate.id == template_id,
            MarketplaceDocumentTemplate.company_id == company_id,
            MarketplaceDocumentTemplate.document_type == 'catalog',
        )
        .first()
    )
    if template is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Marketplace template not found.')
    return template


@router.get('/templates', response_model=CatalogTemplateListResponse)
def list_catalog_templates(
    db: DbSession,
    current_user: ReadUser,
    limit: int = Query(default=50, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
) -> CatalogTemplateListResponse:
    query = db.query(CatalogTemplate).filter(CatalogTemplate.company_id == current_user.company_id)
    total = query.count()
    rows = query.order_by(CatalogTemplate.updated_at.desc()).offset(offset).limit(limit).all()
    return CatalogTemplateListResponse(items=[_to_catalog_template_response(row) for row in rows], total=total)


@router.post('/templates', response_model=CatalogTemplateResponse, status_code=status.HTTP_201_CREATED)
def create_catalog_template(
    payload: CatalogTemplateCreateRequest,
    db: DbSession,
    current_user: WriteUser,
) -> CatalogTemplateResponse:
    existing = (
        db.query(CatalogTemplate)
        .filter(
            CatalogTemplate.company_id == current_user.company_id,
            func.lower(CatalogTemplate.name) == payload.name.strip().lower(),
        )
        .first()
    )
    if existing is not None:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail='Template name already exists.')

    record = CatalogTemplate(
        id=str(uuid4()),
        company_id=current_user.company_id,
        name=payload.name.strip(),
        description=payload.description.strip() if payload.description else None,
        defaults_json=_json_dumps(
            _merge_template_defaults_with_allowed(
                payload.defaults,
                allowed_categories=payload.allowed_categories,
                allowed_style_names=payload.allowed_style_names,
                allowed_compositions=payload.allowed_compositions,
                allowed_woven_knits=payload.allowed_woven_knits,
            )
        ),
        allowed_colors_json=json.dumps(_normalize_template_tokens(payload.allowed_colors)),
        allowed_fabrics_json=json.dumps(_normalize_template_tokens(payload.allowed_fabrics)),
        style_code_pattern=payload.style_code_pattern.strip() if payload.style_code_pattern else None,
        is_active=payload.is_active,
        created_by_user_id=current_user.id,
    )
    db.add(record)
    log_audit(
        db,
        action='catalog.template.create',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'template_id': record.id, 'name': record.name},
    )
    db.commit()
    db.refresh(record)
    return _to_catalog_template_response(record)


@router.patch('/templates/{template_id}', response_model=CatalogTemplateResponse)
def update_catalog_template(
    template_id: str,
    payload: CatalogTemplateUpdateRequest,
    db: DbSession,
    current_user: WriteUser,
) -> CatalogTemplateResponse:
    record = _get_company_template_or_404(db, current_user.company_id, template_id)
    updates = payload.model_dump(exclude_unset=True)
    current_defaults = _json_loads(record.defaults_json)
    current_allowed = _template_allowed_from_defaults(current_defaults)

    if 'name' in updates and updates['name'] is not None:
        normalized_name = updates['name'].strip()
        existing = (
            db.query(CatalogTemplate)
            .filter(
                CatalogTemplate.company_id == current_user.company_id,
                func.lower(CatalogTemplate.name) == normalized_name.lower(),
                CatalogTemplate.id != template_id,
            )
            .first()
        )
        if existing is not None:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail='Template name already exists.')
        record.name = normalized_name

    if 'description' in updates:
        description_value = updates['description']
        record.description = description_value.strip() if isinstance(description_value, str) and description_value.strip() else None
    if (
        'defaults' in updates
        or 'allowed_categories' in updates
        or 'allowed_style_names' in updates
        or 'allowed_compositions' in updates
        or 'allowed_woven_knits' in updates
    ):
        defaults_payload = updates['defaults'] if 'defaults' in updates else _template_defaults_without_meta(current_defaults)
        record.defaults_json = _json_dumps(
            _merge_template_defaults_with_allowed(
                defaults_payload,
                allowed_categories=updates['allowed_categories'] if 'allowed_categories' in updates else current_allowed['categories'],
                allowed_style_names=updates['allowed_style_names'] if 'allowed_style_names' in updates else current_allowed['style_names'],
                allowed_compositions=updates['allowed_compositions'] if 'allowed_compositions' in updates else current_allowed['compositions'],
                allowed_woven_knits=updates['allowed_woven_knits'] if 'allowed_woven_knits' in updates else current_allowed['woven_knits'],
            )
        )
    if 'allowed_colors' in updates:
        record.allowed_colors_json = json.dumps(_normalize_template_tokens(updates['allowed_colors']))
    if 'allowed_fabrics' in updates:
        record.allowed_fabrics_json = json.dumps(_normalize_template_tokens(updates['allowed_fabrics']))
    if 'style_code_pattern' in updates:
        style_pattern = updates['style_code_pattern']
        record.style_code_pattern = style_pattern.strip() if isinstance(style_pattern, str) and style_pattern.strip() else None
    if 'is_active' in updates and updates['is_active'] is not None:
        record.is_active = bool(updates['is_active'])

    log_audit(
        db,
        action='catalog.template.update',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'template_id': record.id, 'name': record.name},
    )
    db.commit()
    db.refresh(record)
    return _to_catalog_template_response(record)


@router.delete('/templates/{template_id}', status_code=status.HTTP_204_NO_CONTENT)
def delete_catalog_template(
    template_id: str,
    db: DbSession,
    current_user: WriteUser,
) -> None:
    record = _get_company_template_or_404(db, current_user.company_id, template_id)
    log_audit(
        db,
        action='catalog.template.delete',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'template_id': record.id, 'name': record.name},
    )
    db.delete(record)
    db.commit()


@router.get('/products', response_model=ProductListResponse)
def list_products(
    db: DbSession,
    current_user: ReadUser,
    search: str | None = Query(default=None, max_length=120),
    status_filter: ProductStatus | None = Query(default=None, alias='status'),
    limit: int = Query(default=25, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
) -> ProductListResponse:
    query = db.query(Product).filter(Product.company_id == current_user.company_id)
    if search:
        term = f'%{search.strip()}%'
        query = query.filter(
            or_(Product.sku.ilike(term), Product.title.ilike(term), Product.category.ilike(term))
        )
    if status_filter:
        query = query.filter(Product.status == status_filter)

    total = query.count()
    rows = query.order_by(Product.updated_at.desc()).offset(offset).limit(limit).all()
    image_map = _build_primary_image_map(db, current_user.company_id, [row.id for row in rows])

    return ProductListResponse(
        items=[_to_product_response(row, primary_image_url=image_map.get(row.id)) for row in rows],
        total=total,
    )


@router.post('/products', response_model=ProductResponse, status_code=status.HTTP_201_CREATED)
def create_product(payload: ProductCreateRequest, db: DbSession, current_user: WriteUser) -> ProductResponse:
    company_settings = (
        db.query(CompanySettings).filter(CompanySettings.company_id == current_user.company_id).first()
    )

    if payload.sku:
        sku = _normalize_explicit_sku(payload.sku)
        existing = (
            db.query(Product)
            .filter(Product.company_id == current_user.company_id, Product.sku == sku)
            .first()
        )
        if existing is not None:
            match = re.match(r'^([A-Z]+[0-9]{2})(\d+)$', sku)
            if match:
                sku = _find_formatted_unique_sku(db, current_user.company_id, match.group(1))
            else:
                sku = _find_unique_sku(db, current_user.company_id, sku)
    else:
        sku = _generate_sku(db, current_user.company_id, company_settings, payload)

    product = Product(
        id=str(uuid4()),
        company_id=current_user.company_id,
        sku=sku,
        title=payload.title.strip(),
        description=payload.description,
        brand=payload.brand,
        category=payload.category,
        color=payload.color,
        size=payload.size,
        status=payload.status,
        confidence_score=payload.confidence_score,
        ai_attributes_json=_json_dumps(payload.ai_attributes),
        created_by_user_id=current_user.id,
    )
    db.add(product)

    log_audit(
        db,
        action='catalog.product.create',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'product_id': product.id, 'sku': product.sku},
    )
    db.commit()
    db.refresh(product)
    return _to_product_response(product)


@router.get('/products/{product_id}', response_model=ProductResponse)
def get_product(product_id: str, db: DbSession, current_user: ReadUser) -> ProductResponse:
    product = _get_company_product_or_404(db, current_user.company_id, product_id)
    image_map = _build_primary_image_map(db, current_user.company_id, [product.id])
    return _to_product_response(product, primary_image_url=image_map.get(product.id))


@router.patch('/products/{product_id}', response_model=ProductResponse)
def update_product(
    product_id: str,
    payload: ProductUpdateRequest,
    db: DbSession,
    current_user: WriteUser,
) -> ProductResponse:
    product = _get_company_product_or_404(db, current_user.company_id, product_id)

    updates = payload.model_dump(exclude_unset=True)
    if 'sku' in updates and updates['sku'] is not None:
        normalized_sku = _normalize_explicit_sku(updates['sku'])
        updates['sku'] = _find_unique_sku(
            db, current_user.company_id, normalized_sku, current_product_id=product.id
        )

    if 'title' in updates and updates['title'] is not None:
        updates['title'] = updates['title'].strip()

    if 'ai_attributes' in updates:
        product.ai_attributes_json = _json_dumps(updates.pop('ai_attributes'))

    for field_name, value in updates.items():
        setattr(product, field_name, value)

    log_audit(
        db,
        action='catalog.product.update',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'product_id': product.id},
    )
    db.commit()
    db.refresh(product)
    image_map = _build_primary_image_map(db, current_user.company_id, [product.id])
    return _to_product_response(product, primary_image_url=image_map.get(product.id))


@router.delete('/products/{product_id}', status_code=status.HTTP_204_NO_CONTENT)
def delete_product(
    product_id: str,
    db: DbSession,
    current_user: WriteUser,
) -> None:
    product = _get_company_product_or_404(db, current_user.company_id, product_id)
    
    log_audit(
        db,
        action='catalog.product.delete',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'product_id': product.id, 'sku': product.sku},
    )
    
    # Delete associated images, measurements, etc. (CASCADE should handle this, but explicit is better)
    db.query(ProductImage).filter(ProductImage.product_id == product_id).delete()
    db.query(ProductMeasurement).filter(ProductMeasurement.product_id == product_id).delete()
    db.query(ProcessingJob).filter(ProcessingJob.product_id == product_id).delete()
    
    db.delete(product)
    db.commit()


@router.post('/products/{product_id}/images', response_model=ProductImageResponse, status_code=status.HTTP_201_CREATED)
def add_product_image(
    product_id: str,
    payload: ProductImageCreateRequest,
    db: DbSession,
    current_user: WriteUser,
) -> ProductImageResponse:
    _get_company_product_or_404(db, current_user.company_id, product_id)
    image = ProductImage(
        id=str(uuid4()),
        company_id=current_user.company_id,
        product_id=product_id,
        file_name=payload.file_name,
        file_url=payload.file_url,
        mime_type=payload.mime_type,
        file_size_bytes=payload.file_size_bytes,
        width_px=payload.width_px,
        height_px=payload.height_px,
        processing_status=payload.processing_status,
        analysis_json=_json_dumps(payload.analysis),
    )
    db.add(image)

    log_audit(
        db,
        action='catalog.product.image.add',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'product_id': product_id, 'image_id': image.id},
    )
    db.commit()
    db.refresh(image)
    return _to_image_response(image)


@router.post(
    '/products/{product_id}/measurements',
    response_model=ProductMeasurementResponse,
    status_code=status.HTTP_201_CREATED,
)
def add_product_measurement(
    product_id: str,
    payload: ProductMeasurementCreateRequest,
    db: DbSession,
    current_user: WriteUser,
) -> ProductMeasurementResponse:
    _get_company_product_or_404(db, current_user.company_id, product_id)
    measurement = ProductMeasurement(
        id=str(uuid4()),
        company_id=current_user.company_id,
        product_id=product_id,
        measurement_key=payload.measurement_key,
        measurement_value=payload.measurement_value,
        unit=payload.unit,
        source=payload.source,
        confidence_score=payload.confidence_score,
        needs_review=payload.needs_review,
        notes=payload.notes,
    )
    db.add(measurement)

    log_audit(
        db,
        action='catalog.product.measurement.add',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'product_id': product_id, 'measurement_id': measurement.id},
    )
    db.commit()
    db.refresh(measurement)
    return _to_measurement_response(measurement)


@router.post('/analyze-image', response_model=AnalyzeImageResponse)
def analyze_image_direct(
    payload: AnalyzeImageRequest,
    db: DbSession,
    current_user: WriteUser,
) -> AnalyzeImageResponse:
    # Import locally or at module level (already imported ai_service at module if we check ai.py, but actually catalog imports process_image_analysis_job not ai_service)
    from app.services.ai import ai_service

    analysis_image_url = _coerce_local_image_url_to_data_url(payload.image_url)
    image_hash = _compute_image_hash(analysis_image_url)
    allowed_options = _normalize_analyze_allowed_options(payload.template_allowed)
    correction_hints = _build_recent_correction_hints(db, current_user.company_id)
    # Synchronously call the AI service and normalize output for UI reliability.
    result = ai_service.analyze_image(
        analysis_image_url,
        allowed_options=allowed_options if allowed_options else None,
        correction_hints=correction_hints if correction_hints else None,
    )
    if isinstance(result, dict) and isinstance(result.get('error'), str):
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f'AI analysis failed: {result["error"]}',
        )
    return _normalize_analysis_result(result, image_hash)


@router.post('/generate-style-code', response_model=GenerateStyleCodeResponse)
def generate_style_code(
    payload: GenerateStyleCodeRequest,
    db: DbSession,
    current_user: WriteUser,
) -> GenerateStyleCodeResponse:
    pattern = payload.pattern.strip() if payload.pattern else ''

    if pattern:
        sku = _render_pattern_sku(db, current_user.company_id, pattern, payload)
        return GenerateStyleCodeResponse(style_code=sku)

    sku = _generate_default_catalog_style_no(db, current_user.company_id, payload.category)
    return GenerateStyleCodeResponse(style_code=sku)


@router.post('/log-correction', response_model=AICorrectionResponse, status_code=status.HTTP_201_CREATED)
def log_correction(
    payload: LogCorrectionRequest,
    background_tasks: BackgroundTasks,
    db: DbSession,
    current_user: WriteUser,
) -> AICorrectionResponse:
    if payload.product_id is not None:
        _get_company_product_or_404(db, current_user.company_id, payload.product_id)

    correction = AICorrection(
        id=str(uuid4()),
        company_id=current_user.company_id,
        product_id=payload.product_id,
        image_hash=payload.image_hash,
        field_name=payload.field_name.strip().lower(),
        feedback_type=payload.feedback_type,
        suggested_value=payload.suggested_value,
        corrected_value=payload.corrected_value,
        reason_code=payload.reason_code,
        notes=payload.notes,
        source=payload.source,
        based_on=payload.based_on,
        learned_from=payload.learned_from,
        confidence_score=payload.confidence_score,
        retraining_status='queued',
        created_by_user_id=current_user.id,
    )
    db.add(correction)
    log_audit(
        db,
        action='catalog.ai_correction.log',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'field_name': correction.field_name, 'feedback_type': correction.feedback_type},
    )
    db.commit()
    db.refresh(correction)

    # Queue lightweight retraining preparation in the background.
    background_tasks.add_task(process_ai_correction_retraining, correction.id)
    return _to_ai_correction_response(correction)


@router.post('/image-labels', response_model=ImageLabelResponse, status_code=status.HTTP_201_CREATED)
def create_image_label(
    payload: ImageLabelCreateRequest,
    db: DbSession,
    current_user: WriteUser,
) -> ImageLabelResponse:
    ai_category = _normalize_optional_label_text(payload.ai_category)
    ai_style = _normalize_optional_label_text(payload.ai_style)
    human_category = _normalize_optional_label_text(payload.human_category) or ai_category
    human_style = _normalize_optional_label_text(payload.human_style) or ai_style
    corrected = bool(payload.corrected or _compute_label_corrected(ai_category, ai_style, human_category, human_style))

    record = ImageLabel(
        id=str(uuid4()),
        company_id=current_user.company_id,
        image_url=payload.image_url.strip(),
        ai_category=ai_category,
        ai_style=ai_style,
        human_category=human_category,
        human_style=human_style,
        corrected=corrected,
    )
    db.add(record)
    log_audit(
        db,
        action='catalog.image_label.create',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'image_label_id': record.id, 'corrected': corrected},
    )
    db.commit()
    db.refresh(record)
    return _to_image_label_response(record)


@router.patch('/image-labels/{label_id}', response_model=ImageLabelResponse)
def update_image_label(
    label_id: str,
    payload: ImageLabelUpdateRequest,
    db: DbSession,
    current_user: WriteUser,
) -> ImageLabelResponse:
    record = (
        db.query(ImageLabel)
        .filter(ImageLabel.id == label_id, ImageLabel.company_id == current_user.company_id)
        .first()
    )
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Image label not found.')

    updates = payload.model_dump(exclude_unset=True)
    if 'human_category' in updates:
        record.human_category = _normalize_optional_label_text(updates['human_category'])
    if 'human_style' in updates:
        record.human_style = _normalize_optional_label_text(updates['human_style'])

    if 'corrected' in updates and updates['corrected'] is not None:
        record.corrected = bool(updates['corrected'])
    else:
        record.corrected = _compute_label_corrected(
            record.ai_category,
            record.ai_style,
            record.human_category,
            record.human_style,
        )

    log_audit(
        db,
        action='catalog.image_label.update',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'image_label_id': record.id, 'corrected': record.corrected},
    )
    db.commit()
    db.refresh(record)
    return _to_image_label_response(record)


@router.get('/image-labels', response_model=ImageLabelListResponse)
def list_image_labels(
    db: DbSession,
    current_user: ReadUser,
    corrected: bool | None = Query(default=None),
    limit: int = Query(default=200, ge=1, le=1000),
) -> ImageLabelListResponse:
    query = db.query(ImageLabel).filter(ImageLabel.company_id == current_user.company_id)
    if corrected is not None:
        query = query.filter(ImageLabel.corrected == corrected)

    total = query.count()
    records = query.order_by(ImageLabel.created_at.desc()).limit(limit).all()
    return ImageLabelListResponse(items=[_to_image_label_response(record) for record in records], total=total)


@router.get('/learning-stats', response_model=LearningStatsResponse)
def get_learning_stats(
    db: DbSession,
    current_user: ReadUser,
) -> LearningStatsResponse:
    items_processed = (
        db.query(func.count(ProductImage.id))
        .filter(
            ProductImage.company_id == current_user.company_id,
            ProductImage.analysis_json.isnot(None),
            ProductImage.analysis_json != '{}',
        )
        .scalar()
        or 0
    )
    corrections = (
        db.query(AICorrection)
        .filter(AICorrection.company_id == current_user.company_id)
        .order_by(AICorrection.created_at.desc())
        .all()
    )
    corrections_received = len(corrections)
    pending_retraining = sum(1 for correction in corrections if correction.retraining_status in {'queued', 'processing'})

    grouped: dict[str, dict[str, int]] = {}
    for correction in corrections:
        field = correction.field_name or 'unknown'
        slot = grouped.setdefault(field, {'accept': 0, 'reject': 0})
        if correction.feedback_type == 'accept':
            slot['accept'] += 1
        else:
            slot['reject'] += 1

    field_accuracy: list[LearningFieldAccuracy] = []
    for field_name, counts in grouped.items():
        accepted = counts['accept']
        rejected = counts['reject']
        total = accepted + rejected
        accuracy = (accepted / total * 100.0) if total > 0 else 0.0
        field_accuracy.append(
            LearningFieldAccuracy(
                field_name=field_name,
                accepted_count=accepted,
                rejected_count=rejected,
                total_feedback=total,
                accuracy_percent=round(accuracy, 2),
            )
        )
    field_accuracy.sort(key=lambda item: item.total_feedback, reverse=True)

    # Conservative estimate for AI-assisted workflow:
    # ~2 min saved per analyzed item + ~30 sec per logged correction.
    time_saved_minutes = int(round((items_processed * 2.0) + (corrections_received * 0.5)))
    insights = _build_learning_insights(items_processed, field_accuracy)

    return LearningStatsResponse(
        items_processed=items_processed,
        corrections_received=corrections_received,
        time_saved_minutes=time_saved_minutes,
        pending_retraining=pending_retraining,
        field_accuracy=field_accuracy,
        insights=insights,
    )


@router.post('/exports', response_model=MarketplaceExportResponse, status_code=status.HTTP_201_CREATED)
def create_export(
    payload: MarketplaceExportCreateRequest,
    db: DbSession,
    current_user: WriteUser,
) -> MarketplaceExportResponse:
    selected_template: MarketplaceDocumentTemplate | None = None
    if payload.template_id:
        selected_template = _get_marketplace_document_template_or_404(db, current_user.company_id, payload.template_id)
        marketplace_key = selected_template.marketplace_key
        template_config = _catalog_template_config_from_record(selected_template)
    else:
        if not payload.marketplace:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail='Marketplace is required when no template is selected.',
            )
        marketplace_key = _normalize_marketplace(payload.marketplace)
        template_config = _catalog_builtin_template_config(marketplace_key)
    normalized_filters = _normalize_export_filters(payload.filters)

    export = MarketplaceExport(
        id=str(uuid4()),
        company_id=current_user.company_id,
        requested_by_user_id=current_user.id,
        marketplace=template_config['label'],
        template_id=selected_template.id if selected_template else None,
        template_name=selected_template.name if selected_template else None,
        export_format=payload.export_format,
        status='processing',
        filters_json=_json_dumps(normalized_filters),
    )
    db.add(export)
    db.flush()

    try:
        products = _query_export_products(db, current_user.company_id, normalized_filters)
        if len(products) == 0:
            raise ValueError('No products matched the selected filters for export.')
        product_ids = [product.id for product in products]
        primary_image_map = _build_primary_image_map(db, current_user.company_id, product_ids)

        validation_errors = _collect_export_validation_errors(products, template_config, primary_image_map)
        if validation_errors:
            preview = '; '.join(validation_errors[:MAX_EXPORT_VALIDATION_ERRORS])
            overflow = len(validation_errors) - MAX_EXPORT_VALIDATION_ERRORS
            suffix = f' (+{overflow} more)' if overflow > 0 else ''
            raise ValueError(f'Validation failed: {preview}{suffix}')

        file_content, row_count = _build_export_file(
            products=products,
            template_config=template_config,
            export_format=payload.export_format,
            primary_image_map=primary_image_map,
        )
        export.file_url = _write_export_file(file_content, marketplace_key, payload.export_format)
        export.row_count = row_count
        export.status = 'completed'
        export.error_message = None
        export.completed_at = datetime.now(timezone.utc)
    except ValueError as err:
        export.status = 'failed'
        export.error_message = str(err)[:512]
        export.file_url = None
        export.row_count = 0
        export.completed_at = datetime.now(timezone.utc)
    except Exception:
        export.status = 'failed'
        export.error_message = 'Export generation failed due to an internal error.'
        export.file_url = None
        export.row_count = 0
        export.completed_at = datetime.now(timezone.utc)

    log_audit(
        db,
        action='catalog.export.create',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={
            'export_id': export.id,
            'marketplace': export.marketplace,
            'status': export.status,
            'row_count': export.row_count,
        },
    )
    db.commit()
    db.refresh(export)
    return _to_export_response(export)


@router.get('/exports', response_model=MarketplaceExportListResponse)
def list_exports(
    db: DbSession,
    current_user: ReadUser,
    marketplace: str | None = Query(default=None, max_length=64),
    status_filter: str | None = Query(default=None, alias='status', max_length=32),
    limit: int = Query(default=25, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
) -> MarketplaceExportListResponse:
    query = db.query(MarketplaceExport).filter(MarketplaceExport.company_id == current_user.company_id)

    if marketplace:
        marketplace_key = _normalize_marketplace(marketplace)
        query = query.filter(MarketplaceExport.marketplace == MARKETPLACE_EXPORT_TEMPLATES[marketplace_key]['label'])

    if status_filter:
        cleaned_status = status_filter.strip().lower()
        if cleaned_status not in ALLOWED_EXPORT_STATUSES:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f'Unsupported export status filter "{status_filter}".',
            )
        query = query.filter(MarketplaceExport.status == cleaned_status)

    total = query.count()
    rows = query.order_by(MarketplaceExport.created_at.desc()).offset(offset).limit(limit).all()
    return MarketplaceExportListResponse(items=[_to_export_response(row) for row in rows], total=total)


@router.post('/jobs', response_model=ProcessingJobResponse, status_code=status.HTTP_201_CREATED)
def create_processing_job(
    payload: ProcessingJobCreateRequest,
    db: DbSession,
    current_user: WriteUser,
    background_tasks: BackgroundTasks,
) -> ProcessingJobResponse:
    if payload.product_id is not None:
        _get_company_product_or_404(db, current_user.company_id, payload.product_id)

    job = ProcessingJob(
        id=str(uuid4()),
        company_id=current_user.company_id,
        product_id=payload.product_id,
        job_type=payload.job_type,
        status='queued',
        input_ref=payload.input_ref,
        payload_json=_json_dumps(payload.payload),
        created_by_user_id=current_user.id,
    )
    db.add(job)

    log_audit(
        db,
        action='catalog.job.create',
        user_id=current_user.id,
        company_id=current_user.company_id,
        metadata={'job_id': job.id, 'job_type': job.job_type},
    )
    db.commit()
    db.refresh(job)
    db.refresh(job)

    if job.job_type == 'image_analysis':
        background_tasks.add_task(process_image_analysis_job, job.id)
    elif job.job_type == 'techpack_ocr':
        background_tasks.add_task(process_techpack_ocr_job, job.id)

    return _to_job_response(job)


@router.get('/jobs', response_model=ProcessingJobListResponse)
def list_processing_jobs(
    db: DbSession,
    current_user: ReadUser,
    status_filter: JobStatus | None = Query(default=None, alias='status'),
    job_type: JobType | None = Query(default=None),
    limit: int = Query(default=25, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
) -> ProcessingJobListResponse:
    query = db.query(ProcessingJob).filter(ProcessingJob.company_id == current_user.company_id)
    if status_filter:
        query = query.filter(ProcessingJob.status == status_filter)
    if job_type:
        query = query.filter(ProcessingJob.job_type == job_type)

    total = query.count()
    rows = query.order_by(ProcessingJob.created_at.desc()).offset(offset).limit(limit).all()
    return ProcessingJobListResponse(items=[_to_job_response(row) for row in rows], total=total)


@router.get('/jobs/{job_id}', response_model=ProcessingJobResponse)
def get_processing_job(
    job_id: str,
    db: DbSession,
    current_user: ReadUser,
) -> ProcessingJobResponse:
    job = db.query(ProcessingJob).filter(ProcessingJob.id == job_id, ProcessingJob.company_id == current_user.company_id).first()
    if not job:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Job not found.')
    return _to_job_response(job)
